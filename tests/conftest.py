from pathlib import Path

import pytest
from docx import Document


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
def docs_root(tmp_path: Path) -> Path:
    """A small document folder, plus a 'secret' file next to it that must never be reachable."""
    root = tmp_path / "docs"
    (root / "policies").mkdir(parents=True)
    (root / ".hidden").mkdir()

    (root / "welcome.txt").write_text("Welcome to the public archive.\nOpening hours are 9-17.\n", encoding="utf-8")
    (root / "notes.md").write_text("# Notes\n\nThe budget meeting is on Friday.\n", encoding="utf-8")
    # Older Windows tools save Swedish text as cp1252, not UTF-8.
    (root / "legacy.txt").write_bytes("Årsredovisning för föreningen".encode("cp1252"))
    (root / "image.png").write_bytes(b"\x89PNG not a document")
    (root / ".hidden" / "draft.txt").write_text("hidden draft", encoding="utf-8")

    doc = Document()
    doc.add_paragraph("Travel policy")
    doc.add_paragraph("Employees should book economy class for trips under six hours.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Per diem"
    table.rows[0].cells[1].text = "450 SEK"
    doc.save(root / "policies" / "travel.docx")

    (tmp_path / "secret.txt").write_text("TOP SECRET budget", encoding="utf-8")
    return root
